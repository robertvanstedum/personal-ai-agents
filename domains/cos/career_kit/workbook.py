"""Career Narrative Kit workbook generator (spec: docs/specs/ career narrative kit, CoS early capability).

Intent and tools are public; data is private. This script writes the blank-to-fill authoring workbook: a landscape
Word table, one row per job (most recent first), project rows, then the About/summary row last. Column 1 carries the
current LinkedIn and resume text; columns 2 (owner notes), 3 (rewrite) and 4 (owner final) start empty. After the main
table come a speaking-points table and a LinkedIn/resume agreement-check table.

The workbook is for authoring only. Resumes sent to employers are separate, single-column documents.

    python -m domains.cos.career_kit.workbook INPUT.json OUTPUT.docx [--carry-from OLD.docx]

A row may carry "rewrite": [lines] for column 3. --carry-from keeps the owner's columns 2–4, speaking points and
agreement decisions from an earlier workbook, by row id; it refuses if a row with owner text would disappear.

It never overwrites: if OUTPUT exists, it refuses (pass a new versioned name). Input format: see example_input.json.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SCHEMA = "career-narrative-workbook/1"
ROW_ID = r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}"         # one rule for the input, the printed [id] and every reader
_ID_TAIL = re.compile(r"\[(" + ROW_ID + r")\]\s*$")
HEADERS = ("Job", "1. Current text (LinkedIn / resume)", "2. Owner notes", "3. Rewrite (AI v0)", "4. Owner final")
WIDTHS = (0.14, 0.28, 0.20, 0.20, 0.18)                    # of the usable page width
RULES = ("Writing rules: change line first, when there was a change · situation, then operating style, then proof · "
         "concrete nouns, not \"transformation\" · LinkedIn text is the entry line; resume adds 2–3 bullets · wins on "
         "the page, the rest as speaking points · no new facts in the rewrite (only what columns 1 and 2 hold) · "
         "nothing confidential unless the owner clears it · owner's voice wins.")


class InputInvalid(ValueError):
    pass


def _text(value, field):
    if value is None:
        return ""
    if not isinstance(value, str):
        raise InputInvalid(f"{field} must be text")
    return value.strip()


def load(path: Path) -> dict:
    doc = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(doc, dict) or doc.get("schema") != SCHEMA:
        raise InputInvalid(f"schema must be {SCHEMA!r}")
    rows = doc.get("rows")
    if not isinstance(rows, list) or not rows:
        raise InputInvalid("rows must be a non-empty list")
    ids = set()
    for i, row in enumerate(rows):
        if not isinstance(row, dict) or row.get("kind") not in ("job", "project", "about"):
            raise InputInvalid(f"rows[{i}].kind must be job, project or about")
        rid = _text(row.get("id"), f"rows[{i}].id")
        if not rid or rid in ids:
            raise InputInvalid(f"rows[{i}].id missing or duplicate")
        if not re.fullmatch(ROW_ID, rid):
            raise InputInvalid(f"rows[{i}].id {rid!r}: letters, digits, _ . - only (it is printed in the workbook as [id])")
        ids.add(rid)
        bullets = row.get("resume_bullets", [])
        if not isinstance(bullets, list) or not all(isinstance(b, str) for b in bullets):
            raise InputInvalid(f"rows[{i}].resume_bullets must be a list of text")
    if rows[-1]["kind"] != "about" or sum(r["kind"] == "about" for r in rows) != 1:
        raise InputInvalid("exactly one about row, and it comes last")
    for i, item in enumerate(doc.get("agreement", [])):
        if not isinstance(item, dict) or not _text(item.get("row"), f"agreement[{i}].row"):
            raise InputInvalid(f"agreement[{i}] needs a row")
    return doc


def _shade(cell, hex_fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def _repeat_header(row):
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def _para(cell, text, *, bold=False, size=9, label=None, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if label:
        r = p.add_run(label + " ")
        r.bold, r.font.size = True, Pt(size)
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def _table(document, headers, widths, usable):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, col in enumerate(table.columns):            # the grid, which viewers other than Word read
        col.width = Inches(usable * widths[i])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(usable * widths[i])
        _para(cell, h, bold=True, first=True)
        _shade(cell, "D9E2F3")
    _repeat_header(table.rows[0])
    return table


def _row(table, widths, usable):
    cells = table.add_row().cells
    for i, c in enumerate(cells):
        c.width = Inches(usable * widths[i])
    return cells


def _job_cell(cell, row):
    kind = {"job": "", "project": "Project · ", "about": ""}[row["kind"]]
    if row["kind"] == "about":
        _para(cell, "About / summary, headline (write last)", bold=True, first=True)
        return
    _para(cell, kind + _text(row.get("title"), "title"), bold=True, first=True)
    for field in ("employer", "dates", "place"):
        value = _text(row.get(field), field)
        if value:
            _para(cell, value)
    _para(cell, f"[{row['id']}]", size=7)


def _current_cell(cell, row):
    linkedin = _text(row.get("linkedin"), "linkedin") or "(not on LinkedIn)"
    _para(cell, linkedin, label="LinkedIn:", first=True)
    entry = _text(row.get("resume_entry"), "resume_entry")
    if entry or row.get("resume_bullets"):
        _para(cell, entry, label="Resume:")
        for b in row.get("resume_bullets", []):
            _para(cell, "• " + b.strip())
    else:
        _para(cell, "(not on the resume)", label="Resume:")


def read_carry(old: Path) -> dict:
    """What the owner wrote in an earlier workbook: columns 2–4 by row id, speaking points by row id, agreement
    decisions by label. Word list bullets come back as "• …"."""
    from .export import cell_lines, row_id                 # lazy: export imports this module
    tables = Document(old).tables
    carry = {"cols": {}, "speaking": {}, "agreement": {}}
    for row in tables[0].rows[1:]:
        rid = row_id(row.cells[0])
        if rid:
            carry["cols"][rid] = {i: cell_lines(row.cells[i]) for i in (2, 3, 4)}
    if len(tables) > 1:
        for row in tables[1].rows[1:]:
            m = _ID_TAIL.search(row.cells[0].text.strip())
            lines = [p.text.strip() for p in row.cells[1].paragraphs if p.text.strip()]
            if m and lines:
                carry["speaking"][m.group(1)] = lines
    if len(tables) > 2:
        for row in tables[2].rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if cells and cells[0]:
                carry["agreement"][cells[0]] = cells[1:4] + [""] * (3 - len(cells[1:4]))
    return carry


def _lines_into(cell, lines):
    for i, line in enumerate(lines):
        _para(cell, line, first=(i == 0))


def build(doc: dict, out: Path, carry_from: Path | None = None) -> Path:
    """carry_from: an earlier workbook whose columns 2–4, speaking points and decisions are carried forward by row id
    (owner text always wins over the input's rewrite). Rows that no longer exist are reported, never silently lost."""
    if out.exists():
        raise FileExistsError(f"{out} exists; write a new versioned name instead")
    carry = read_carry(carry_from) if carry_from else {"cols": {}, "speaking": {}, "agreement": {}}
    ids = {r["id"] for r in doc["rows"]}
    orphans = sorted({rid for rid, cols in carry["cols"].items() if rid not in ids and any(cols.values())}
                     | {rid for rid in carry["speaking"] if rid not in ids})
    if orphans:
        raise InputInvalid("the earlier workbook has rows this input lacks (their text would be lost): "
                           + ", ".join(orphans))
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.5))
    usable = 10.0
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(9)

    document.add_heading(_text(doc.get("title"), "title") or "Career narrative workbook", level=1)
    meta = " · ".join(x for x in (_text(doc.get("variant"), "variant") and f"Variant: {doc['variant']}",
                                  _text(doc.get("version"), "version") and f"Version: {doc['version']}") if x)
    if meta:
        document.add_paragraph(meta)
    rules = document.add_paragraph(RULES)
    rules.runs[0].font.size = Pt(8)
    rules.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_paragraph("Order of work: most recent job first; the About/summary row last. Fill column 2 before "
                           "anything is written in column 3. Column 4 is what ships.").runs[0].font.size = Pt(8)

    table = _table(document, HEADERS, WIDTHS, usable)
    for row in doc["rows"]:
        cells = _row(table, WIDTHS, usable)
        _job_cell(cells[0], row)
        _current_cell(cells[1], row)
        old = carry["cols"].get(row["id"], {})
        _lines_into(cells[2], old.get(2) or [])
        _lines_into(cells[3], old.get(3) or row.get("rewrite") or [])
        _lines_into(cells[4], old.get(4) or [])

    document.add_heading("Speaking points (off the page)", level=2)
    sp = _table(document, ("Job", "Details kept for interviews, not for the resume"), (0.25, 0.75), usable)
    for row in doc["rows"]:
        if row["kind"] != "about":
            cells = _row(sp, (0.25, 0.75), usable)
            _para(cells[0], f"{_text(row.get('employer'), 'employer') or _text(row.get('title'), 'title')} [{row['id']}]",
                  first=True)
            _lines_into(cells[1], carry["speaking"].get(row["id"])
                        or [_text(row.get("speaking_points"), "speaking_points")])

    document.add_heading("Agreement check (LinkedIn vs resume)", level=2)
    ac = _table(document, ("Job", "LinkedIn", "Resume", "Decision"), (0.18, 0.32, 0.32, 0.18), usable)
    seen = set()
    items = [(_text(i.get("row"), "row"), [_text(i.get("linkedin"), "linkedin"), _text(i.get("resume"), "resume"), ""])
             for i in doc.get("agreement", [])]
    items += [(label, ["", "", ""]) for label in carry["agreement"] if label not in {l for l, _ in items}]
    for label, values in items:                     # owner text wins, cell by cell; owner-added rows are kept
        if label in seen:
            continue
        seen.add(label)
        owner = carry["agreement"].get(label, ["", "", ""])
        cells = _row(ac, (0.18, 0.32, 0.32, 0.18), usable)
        _para(cells[0], label, first=True)
        for i in range(3):
            _para(cells[i + 1], owner[i] or values[i], first=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)
    return out


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    carry, bad = None, False
    if "--carry-from" in argv:
        i = argv.index("--carry-from")
        if i + 1 < len(argv):
            carry = Path(argv[i + 1])
        else:
            bad = True
        del argv[i:i + 2]
    if len(argv) != 2 or bad:
        print("usage: python -m domains.cos.career_kit.workbook INPUT.json OUTPUT.docx [--carry-from OLD.docx]",
              file=sys.stderr)
        return 2
    try:
        print(build(load(Path(argv[0])), Path(argv[1]), carry_from=carry))
    except (InputInvalid, FileExistsError, json.JSONDecodeError) as exc:
        print(f"refused: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
