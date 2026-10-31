"""Read a resume .docx into workbook rows (column 1, resume side).

Expected layout (the kit's resume format): section headings in capitals (PROFESSIONAL SUMMARY, CURRENT WORK,
PROFESSIONAL EXPERIENCE, EDUCATION); each role a line "Title | Employer   Place", then a dates line, then entry
paragraphs, then bullets in the "List Paragraph" style. CURRENT WORK roles become project rows; the summary becomes the
About row. LinkedIn text is not in a resume: the caller adds it (or leaves it empty, shown as "not on LinkedIn").

    python -m domains.cos.career_kit.resume_import RESUME.docx OUTPUT.json [--title T] [--variant V]

Refuses to overwrite OUTPUT. Prints one line per row so the owner can check the split.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document

from .workbook import SCHEMA

SECTIONS = ("PROFESSIONAL SUMMARY", "CURRENT WORK", "PROFESSIONAL EXPERIENCE", "EDUCATION")


class LayoutUnknown(ValueError):
    pass


def _slug(text: str, taken: set) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:24] or "row"
    slug, n = base, 2
    while slug in taken:
        slug, n = f"{base}-{n}", n + 1
    taken.add(slug)
    return slug


def _roles(paras, kind, taken):
    rows, cur = [], None
    for style, text in paras:
        if "|" in text and style != "List Paragraph":
            title, rest = (x.strip() for x in text.split("|", 1))
            parts = re.split(r"\s{3,}", rest, maxsplit=1)
            employer, place = parts[0].strip(), (parts[1].strip() if len(parts) > 1 else "")
            cur = {"id": _slug(employer + " " + place.split(",")[0], taken), "kind": kind, "title": title,
                   "employer": employer, "dates": None, "place": place, "linkedin": None, "resume_entry": "",
                   "resume_bullets": []}
            rows.append(cur)
        elif cur is None:
            continue
        elif cur["dates"] is None:
            cur["dates"] = text
        elif style == "List Paragraph":
            cur["resume_bullets"].append(text)
        else:
            cur["resume_entry"] = (cur["resume_entry"] + " " + text).strip()
    return rows


def read(path: Path) -> list[dict]:
    paras = [(p.style.name, p.text.strip()) for p in Document(path).paragraphs if p.text.strip()]
    heads = {t: i for i, (_, t) in enumerate(paras) if t in SECTIONS}
    if "PROFESSIONAL EXPERIENCE" not in heads or "EDUCATION" not in heads:
        raise LayoutUnknown("needs PROFESSIONAL EXPERIENCE and EDUCATION headings")
    order = sorted(heads.values())

    def section(name):
        if name not in heads:
            return []
        start = heads[name]
        end = next((i for i in order if i > start), len(paras))
        return paras[start + 1:end]

    taken: set = set()
    jobs = _roles(section("PROFESSIONAL EXPERIENCE"), "job", taken)
    projects = _roles(section("CURRENT WORK"), "project", taken)
    if not jobs:
        raise LayoutUnknown("no roles found under PROFESSIONAL EXPERIENCE")
    summary = " ".join(t for _, t in section("PROFESSIONAL SUMMARY"))
    about = {"id": "about", "kind": "about", "linkedin": None,
             "resume_entry": ("Professional summary: " + summary) if summary else "", "resume_bullets": []}
    return jobs + projects + [about]


def read_frame(path: Path) -> tuple[list[str], list[dict]]:
    """What surrounds the workbook rows on the resume, carried through unchanged by the export: the header lines
    before the first section (name, contact) and every paragraph from EDUCATION to the end (education,
    capabilities ...), each as {"text", "bullet"}."""
    paras = [(p.style.name, p.text.strip()) for p in Document(path).paragraphs if p.text.strip()]
    texts = [t for _, t in paras]
    first = min((texts.index(s) for s in SECTIONS if s in texts), default=0)
    tail = texts.index("EDUCATION") if "EDUCATION" in texts else len(paras)
    return texts[:first], [{"text": t, "bullet": s == "List Paragraph"} for s, t in paras[tail:]]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("resume", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--title", default="Career narrative workbook")
    ap.add_argument("--variant", default="")
    args = ap.parse_args(argv)
    if args.output.exists():
        print(f"refused: {args.output} exists; write a new versioned name instead", file=sys.stderr)
        return 1
    try:
        rows = read(args.resume)
    except LayoutUnknown as exc:
        print(f"refused: {exc}", file=sys.stderr)
        return 1
    header, appendix = read_frame(args.resume)
    doc = {"schema": SCHEMA, "title": args.title, "variant": args.variant, "header": header, "rows": rows,
           "appendix": appendix, "agreement": []}
    args.output.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    for r in rows:
        print(f"{r['id']:<28} {r['kind']:<8} {r.get('title') or ''} | {r.get('dates') or ''} | "
              f"{len(r['resume_bullets'])} bullets")
    return 0


if __name__ == "__main__":
    sys.exit(main())
