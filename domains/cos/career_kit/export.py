"""Export a finished workbook: column 4 (owner final) becomes the outward-facing texts.

    python -m domains.cos.career_kit.export INPUT.json WORKBOOK.docx OUTDIR [--draft]

--draft previews from column 3 (the rewrite) instead; every output is then marked DRAFT.
--template RESUME.docx writes the resume into the owner's own resume, keeping its layout, header and Education/
Capabilities sections (recommended: the resume the workbook was imported from).
Paragraphs made bullets with Word's own list button count as bullets.

INPUT.json is the generator's input (titles, employers, dates, places, optional "header" lines for the resume top).
WORKBOOK.docx is the owner's edited workbook; rows are matched by the [id] printed in the Job cell. In column 4:
  - a paragraph starting "LinkedIn:" is the LinkedIn entry text for that role;
  - paragraphs starting "•" or "-" are resume bullets;
  - any other paragraph is the resume entry line;
  - in the About row: "Headline:", "About:" (LinkedIn) and "Summary:" (resume) paragraphs;
  - a paragraph starting "Note:" is an editorial note and is never exported.
A row whose column 4 is empty is NOT exported with AI or old text: it is listed in the report as missing.

Writes into OUTDIR (which must not exist): resume.docx (single column, no tables or text boxes, so applicant tracking
systems read it), linkedin.md, and report.md (missing finals, length flags, LinkedIn/resume agreement flags).
Length limits are LinkedIn's published limits at the time of writing; verify them when they matter.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from .workbook import load

LIMITS = {"headline": 220, "about": 2600, "experience": 2000}       # characters; verify when they matter
from .workbook import _ID_TAIL as _ID  # noqa: E402  (one id rule for writer and reader)


FINAL, REWRITE = 4, 3                                              # workbook columns


def _is_list(p) -> bool:
    """A paragraph the owner made a bullet with Word's own list button (a list style or list numbering)."""
    ppr = p._p.pPr
    return p.style.name.startswith("List") or (ppr is not None and ppr.numPr is not None)


def cell_lines(cell) -> list[str]:
    """The cell's paragraphs as text; Word list paragraphs come back as "• …" so they stay bullets."""
    out = []
    for p in cell.paragraphs:
        text = p.text.strip()
        if text:
            out.append("• " + text if _is_list(p) and text[:1] not in ("•", "-") else text)
    return out


def row_id(cell) -> str | None:
    lines = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    if lines and lines[0].startswith("About / summary"):
        return "about"
    m = _ID.search(lines[-1]) if lines else None
    return m.group(1) if m else None


def read_finals(workbook: Path, column: int = FINAL) -> dict:
    """{row id: [paragraph texts of that column]}, from the main table (column 4, the owner final, by default)."""
    table = Document(workbook).tables[0]
    out = {}
    for row in table.rows[1:]:
        rid = row_id(row.cells[0])
        if rid:
            out[rid] = cell_lines(row.cells[column])
    return out


def read_decisions(workbook: Path) -> dict:
    """{agreement row label: the owner's Decision cell}, from the workbook's agreement-check table."""
    tables = Document(workbook).tables
    if len(tables) < 3:
        return {}
    out = {}
    for row in tables[2].rows[1:]:
        cells = row.cells
        label, decision = cells[0].text.strip(), cells[-1].text.strip()
        if label and decision:
            out[label] = decision
    return out


def parse_final(lines: list[str], about: bool) -> dict:
    final = {"entry": [], "bullets": [], "linkedin": [], "headline": "", "about": [], "summary": []}
    for line in lines:
        low = line.lower()
        if low.startswith("note:"):                           # editorial note between owner and helper
            continue
        if about and low.startswith("headline:"):
            final["headline"] = line.split(":", 1)[1].strip()
        elif about and low.startswith("about:"):
            final["about"].append(line.split(":", 1)[1].strip())
        elif about and low.startswith("summary:"):
            final["summary"].append(line.split(":", 1)[1].strip())
        elif low.startswith("linkedin:"):
            final["linkedin"].append(line.split(":", 1)[1].strip())
        elif line[:1] in ("•", "-"):
            final["bullets"].append(line[1:].strip())
        elif about:
            final["summary"].append(line)
        else:
            final["entry"].append(line)
    return final


def _bold(document, text, size=None):
    run = document.add_paragraph().add_run(text)
    run.bold = True
    if size:
        run.font.size = Pt(size)


class PlainWriter:
    """A clean single-column resume when there is no template; same calls as resume_template.ResumeWriter."""

    def __init__(self, header, appendix):
        self.doc, self.appendix = Document(), appendix
        self.doc.styles["Normal"].font.name, self.doc.styles["Normal"].font.size = "Calibri", Pt(10.5)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(self.doc.sections[0], side, Inches(0.75))
        for i, line in enumerate(header):
            _bold(self.doc, line, 16) if i == 0 else self.doc.add_paragraph(line)

    def heading(self, text):
        _bold(self.doc, text)

    def spacer(self):
        pass

    def role(self, title, employer, place):
        _bold(self.doc, f"{title}  |  {employer}" + (f"   {place}" if place else ""))

    def dates(self, text):
        self.doc.add_paragraph(text)

    def body(self, text):
        self.doc.add_paragraph(text)

    def bullet(self, text):
        self.doc.add_paragraph(text, style="List Bullet")

    def save(self, out):
        for item in self.appendix:                          # education, capabilities ...: carried unchanged
            if item.get("bullet"):
                self.doc.add_paragraph(item["text"], style="List Bullet")
            elif item["text"].isupper():
                _bold(self.doc, item["text"])
            else:
                self.doc.add_paragraph(item["text"])
        self.doc.save(out)


def export(input_path: Path, workbook: Path, outdir: Path, *, draft: bool = False, pages: bool = True,
           template: Path | None = None) -> Path:
    """draft=True previews from column 3 (the rewrite); every output then says DRAFT and nothing is a final.
    template: the owner's own resume .docx; the export is written into it so it keeps the owner's layout."""
    if outdir.exists():
        raise FileExistsError(f"{outdir} exists; use a new versioned folder")
    doc = load(input_path)
    finals = read_finals(workbook, REWRITE if draft else FINAL)
    missing, flags, linkedin_md = [], [], []
    if template is not None:
        from .resume_template import ResumeWriter
        w = ResumeWriter(template)
    else:
        w = PlainWriter(doc.get("header", []), doc.get("appendix", []))
    rows = doc["rows"]
    about = parse_final(finals.get("about", []), about=True)
    if not finals.get("about"):
        missing.append("about")
    if draft:
        w.heading("DRAFT PREVIEW FROM COLUMN 3, NOT THE OWNER'S FINAL TEXT. DO NOT SEND.")
    if about["summary"]:
        w.heading("PROFESSIONAL SUMMARY")
        for s in about["summary"]:
            w.body(s)
        w.spacer()
    if about["headline"]:
        linkedin_md += ["## Headline", "", about["headline"], ""]
        if len(about["headline"]) > LIMITS["headline"]:
            flags.append(f"headline: {len(about['headline'])} characters > {LIMITS['headline']}")
    if about["about"]:
        text = "\n\n".join(about["about"])
        linkedin_md += ["## About", "", text, ""]
        if len(text) > LIMITS["about"]:
            flags.append(f"about: {len(text)} characters > {LIMITS['about']}")
    for kind, heading in (("project", "CURRENT WORK"), ("job", "PROFESSIONAL EXPERIENCE")):
        group = [r for r in rows if r["kind"] == kind]
        if not group:
            continue
        w.heading(heading)
        if kind == "job":
            linkedin_md += ["## Experience", ""]
        for r in group:
            lines = finals.get(r["id"], [])
            if not lines:
                missing.append(r["id"])
                continue
            f = parse_final(lines, about=False)
            w.role(r.get("title", ""), r.get("employer", ""), r.get("place", ""))
            if r.get("dates"):
                w.dates(r["dates"])
            for e in f["entry"]:
                w.body(e)
            for b in f["bullets"]:
                w.bullet(b)
            if kind == "job" and f["linkedin"]:
                text = " ".join(f["linkedin"])
                linkedin_md += [f"### {r.get('title', '')} · {r.get('employer', '')}", "", text, ""]
                if len(text) > LIMITS["experience"]:
                    flags.append(f"{r['id']}: LinkedIn entry {len(text)} characters > {LIMITS['experience']}")
        w.spacer()
    decisions = read_decisions(workbook)
    agreement = [f"- {a['row']}: LinkedIn \"{a.get('linkedin', '')}\" · resume \"{a.get('resume', '')}\""
                 + (f" · decision: {decisions[a['row']]}" if decisions.get(a["row"]) else " · decision open")
                 for a in doc.get("agreement", [])]
    outdir.mkdir(parents=True)
    w.save(outdir / "resume.docx")
    banner = ["DRAFT preview from column 3 (rewrite), not the owner's final text. Do not post.", ""] if draft else []
    (outdir / "linkedin.md").write_text("\n".join(banner + linkedin_md) + "\n", encoding="utf-8")
    from .pages import page_count
    counted = page_count(outdir / "resume.docx") if pages else None
    report = ["# Export report" + (" — DRAFT (column 3)" if draft else ""), "",
              f"Input: {input_path.name} · workbook: {workbook.name} · source: "
              + ("column 3, rewrite (draft preview)" if draft else "column 4, owner final"), "",
              "## Resume length", "",
              (f"- {counted} pages (rendered with LibreOffice; Word may differ slightly)" if counted
               else "- not checked (LibreOffice not installed, or rendering failed)"), "",
              "## Rows without text in the source column (not exported)", ""] + \
             ([f"- {m}" for m in missing] or ["- none"]) + \
             ["", "## Length flags (LinkedIn)", ""] + ([f"- {x}" for x in flags] or ["- none"]) + \
             ["", "## LinkedIn / resume agreement (review; nothing is auto-corrected)", ""] + (agreement or ["- none"])
    (outdir / "report.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    return outdir


def main(argv=None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    argv = list(argv)
    draft = "--draft" in argv
    argv = [a for a in argv if a != "--draft"]
    template = None
    if "--template" in argv:
        i = argv.index("--template")
        template = Path(argv[i + 1]) if i + 1 < len(argv) else None
        del argv[i:i + 2]
        if template is None:
            argv = []
    if len(argv) != 3:
        print("usage: python -m domains.cos.career_kit.export INPUT.json WORKBOOK.docx OUTDIR [--draft] "
              "[--template RESUME.docx]", file=sys.stderr)
        return 2
    try:
        print(export(Path(argv[0]), Path(argv[1]), Path(argv[2]), draft=draft, template=template))
    except (FileExistsError, ValueError, json.JSONDecodeError) as exc:
        print(f"refused: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
